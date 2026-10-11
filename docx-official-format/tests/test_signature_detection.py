#!/usr/bin/env python3
"""Regression tests for official-document signature/date formatting."""

from __future__ import annotations

import importlib.util
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "format_docx.py"


def cjk_font(run) -> str | None:
    """读中文字体槽位。

    `run.font.name` 只读 w:ascii 槽位，现该槽位固定为 Times New Roman，
    校验中文字体必须读 w:eastAsia。
    """
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return None
    rfonts = rpr.find(qn("w:rFonts"))
    return None if rfonts is None else rfonts.get(qn("w:eastAsia"))


def latin_font(run) -> str | None:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return None
    rfonts = rpr.find(qn("w:rFonts"))
    return None if rfonts is None else rfonts.get(qn("w:ascii"))


def hint(run) -> str | None:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return None
    rfonts = rpr.find(qn("w:rFonts"))
    return None if rfonts is None else rfonts.get(qn("w:hint"))


def load_formatter():
    spec = importlib.util.spec_from_file_location("official_docx_formatter", SCRIPT_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load formatter script: {SCRIPT_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


FORMATTER = load_formatter()


def convert_document(paragraphs: list[str]) -> Document:
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        input_path = tmpdir_path / "input.docx"
        output_path = tmpdir_path / "output.docx"

        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(input_path)

        FORMATTER.format_docx(
            input_path,
            output_path,
            title_font="FZXiaoBiaoSong-B05",
            body_font="FangSong_GB2312",
            heading_font="SimHei",
            secondary_heading_font="KaiTi",
        )
        return Document(output_path)


def convert(paragraphs: list[str]) -> list:
    return [p for p in convert_document(paragraphs).paragraphs if p.text.strip()]


def convert_with_report(paragraphs: list[str]):
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        input_path = tmpdir_path / "input.docx"
        output_path = tmpdir_path / "output.docx"

        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(input_path)

        report = FORMATTER.format_docx_with_report(
            input_path,
            output_path,
            title_font="FZXiaoBiaoSong-B05",
            body_font="FangSong_GB2312",
            heading_font="SimHei",
            secondary_heading_font="KaiTi",
            check_typos=True,
        )
        return Document(output_path), report


class SignatureDetectionTest(unittest.TestCase):
    def assert_length_close(self, actual, expected) -> None:
        self.assertLessEqual(abs(int(actual) - int(expected)), 500)

    def assert_tail_is_body(self, paragraphs: list[str]) -> None:
        output_paragraphs = convert(paragraphs)
        for paragraph in output_paragraphs[-2:]:
            self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertIsNotNone(paragraph.paragraph_format.first_line_indent)
            self.assertIsNone(paragraph.paragraph_format.right_indent)

    def assert_tail_is_signature(self, paragraphs: list[str]) -> None:
        output_paragraphs = convert(paragraphs)
        for paragraph in output_paragraphs[-2:]:
            self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            self.assertIsNone(paragraph.paragraph_format.first_line_indent)
            self.assertIsNotNone(paragraph.paragraph_format.right_indent)

    def test_tail_date_without_signature_stays_body(self) -> None:
        self.assert_tail_is_body(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请大家按要求完成报名。",
                "注意事项",
                "2026年5月13日",
            ]
        )

    def test_text_containing_date_stays_body(self) -> None:
        output_paragraphs = convert(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请大家按要求完成报名。",
                "培训时间为2026年5月13日",
            ]
        )
        tail = output_paragraphs[-1]
        self.assertEqual(tail.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertIsNotNone(tail.paragraph_format.first_line_indent)
        self.assertIsNone(tail.paragraph_format.right_indent)

    def test_organization_signature_and_date_align_right(self) -> None:
        self.assert_tail_is_signature(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请大家按要求完成报名。",
                "重庆市科学技术局",
                "2026年5月13日",
            ]
        )

    def test_person_signature_and_date_align_right(self) -> None:
        self.assert_tail_is_signature(
            [
                "个人事项说明",
                "本人已完成相关材料整理。",
                "张军",
                "2026年5月13日",
            ]
        )

    def test_signature_and_date_share_same_right_indent(self) -> None:
        """长署名与成文日期右缩进必须相等（都右空四字）。

        署名比日期宽时，若按「居成文日期中」把署名右缩进调大，
        会出现署名缩进比日期更多、视觉上不在同一右边缘的错位。
        """
        output_paragraphs = convert(
            [
                "“新韵重庆”无人机灯光秀",
                "10月展演安排通告",
                "一、时间安排",
                "2026年10月1日至6日（国庆节）20：30。",
                "“新韵重庆”无人机灯光秀指挥部",
                "2026年9月28日",
            ]
        )
        issuer, date = output_paragraphs[-2], output_paragraphs[-1]
        self.assertEqual(issuer.paragraph_format.right_indent, date.paragraph_format.right_indent)
        # 三号字 16pt，右空四字 = 64pt
        self.assertEqual(int(issuer.paragraph_format.right_indent.pt), 64)

    def test_official_format_title_line_spacing_and_heading_fonts(self) -> None:
        output_paragraphs = convert(
            [
                "重庆市商务委员会",
                "关于开展AI办公培训的通知",
                "",
                "各处室、委属单位：",
                "为提升日常办公效率，现组织开展AI办公培训。",
                "一、工作安排",
                "（一）培训对象。",
                "请各部门按要求组织参加。",
            ]
        )
        by_text = {paragraph.text.strip(): paragraph for paragraph in output_paragraphs}

        self.assertEqual(by_text["重庆市商务委员会"].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(by_text["关于开展AI办公培训的通知"].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(by_text["各处室、委属单位："].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(by_text["为提升日常办公效率，现组织开展AI办公培训。"].paragraph_format.line_spacing, Pt(30))

        primary_run = by_text["一、工作安排"].runs[0]
        self.assertEqual(cjk_font(primary_run), "SimHei")

        secondary_run = by_text["（一）培训对象。"].runs[0]
        self.assertEqual(cjk_font(secondary_run), "KaiTi")

    def test_mixed_script_uses_four_slot_font_assignment(self) -> None:
        """西文与数字走 Times New Roman，汉字走中文字体，全角引号独占 run 且 hint=eastAsia。"""
        document = convert_document(
            [
                "关于开展AI办公培训的通知",
                "各处室、委属单位：",
                "“渝见”培训2026年10月20日（星期一）9：30开班。",
            ]
        )

        target = next(p for p in document.paragraphs if "培训2026" in p.text)

        latin_runs = [r for r in target.runs if any(c.isascii() and c.isalnum() for c in r.text)]
        self.assertTrue(latin_runs, "西文与数字应拆成独立 run")
        for run in latin_runs:
            self.assertEqual(latin_font(run), "Times New Roman")
            self.assertEqual(hint(run), "default")

        quote_runs = [r for r in target.runs if r.text in "“”"]
        self.assertEqual(len(quote_runs), 2, "全角引号应各自独占一个 run")
        for run in quote_runs:
            self.assertEqual(cjk_font(run), "FangSong_GB2312")
            self.assertEqual(hint(run), "eastAsia")

        for run in target.runs:
            self.assertEqual(latin_font(run), "Times New Roman")

    def test_local_print_format_page_margins_are_applied(self) -> None:
        document = convert_document(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请大家按要求完成报名。",
            ]
        )
        section = document.sections[0]
        self.assert_length_close(section.top_margin, Mm(35))
        self.assert_length_close(section.bottom_margin, Mm(29))
        self.assert_length_close(section.left_margin, Mm(25.5))
        self.assert_length_close(section.right_margin, Mm(25.5))

    def test_page_numbers_use_outside_alignment_after_first_page(self) -> None:
        document = convert_document(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请大家按要求完成报名。",
            ]
        )
        self.assertTrue(document.settings.odd_and_even_pages_header_footer)
        section = document.sections[0]
        self.assertTrue(section.different_first_page_header_footer)
        self.assertEqual(section.first_page_footer.paragraphs[0].text, "")
        self.assertEqual(section.footer.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(section.even_page_footer.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(section.footer.paragraphs[0].text, "— 1 —")
        self.assertEqual(section.even_page_footer.paragraphs[0].text, "— 1 —")

    def test_first_page_footer_has_no_field_characters(self) -> None:
        """首页页脚必须连域元素一起清掉。

        只断言 `paragraph.text == ""` 是不够的：把 run.text 置空后，
        run 内的 w:fldChar / w:instrText 仍然存在，域照样生效，
        首页仍会显示「— 1 —」。必须断言 XML 里没有任何域标记。
        """
        document = convert_document(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请大家按要求完成报名。",
            ]
        )
        first_footer = document.sections[0].first_page_footer
        xml = first_footer._element.xml
        for marker in ("w:fldChar", "w:instrText", "PAGE"):
            self.assertNotIn(marker, xml, f"首页页脚仍残留 {marker}")

    def test_page_number_field_uses_separate_runs(self) -> None:
        """页码域必须拆成独立 run，begin/separate/end 不可塞进同一个 run。"""
        document = convert_document(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请大家按要求完成报名。",
            ]
        )
        for footer in (
            document.sections[0].footer,
            document.sections[0].even_page_footer,
        ):
            paragraph = footer.paragraphs[0]

            #每个 fldChar 必须独占一个 run，且 run 内不得混入其他标记
            fld_runs = [
                run for run in paragraph.runs if run._element.find(qn("w:fldChar")) is not None
            ]
            self.assertEqual(len(fld_runs), 3, "begin / separate / end 应各占一个 run")
            for run in fld_runs:
                children = [child.tag.split("}")[1] for child in run._element]
                self.assertEqual(
                    children.count("fldChar"), 1, "一个 run 内不得塞多个 fldChar"
                )
                self.assertNotIn("instrText", children)
                self.assertNotIn("t", children)

            instr_runs = [
                run for run in paragraph.runs if run._element.find(qn("w:instrText")) is not None
            ]
            self.assertEqual(len(instr_runs), 1)
            self.assertEqual(
                instr_runs[0]._element.find(qn("w:instrText")).text.strip(), "PAGE"
            )
            # PAGE 域前后要有缓存结果 run
            self.assertEqual(paragraph.text, "— 1 —")

    def test_punctuation_and_date_are_normalized(self) -> None:
        document, report = convert_with_report(
            [
                "关于开展AI办公培训的通知。",
                "各位老师:",
                "附件:培训安排。",
                "重庆市商务委员会",
                "二〇二六年六月四日",
            ]
        )
        texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        self.assertIn("关于开展AI办公培训的通知", texts)
        self.assertIn("各位老师：", texts)
        self.assertIn("附件：培训安排", texts)
        self.assertIn("2026年6月4日", texts)
        self.assertGreaterEqual(report.punctuation_fix_count, 4)

    def test_punctuation_issue_report_marks_unbalanced_pairs(self) -> None:
        _document, report = convert_with_report(
            [
                "关于开展AI办公培训的通知",
                "各位老师：",
                "请认真学习《标点符号用法。",
            ]
        )
        self.assertGreaterEqual(report.punctuation_issue_count, 1)

    def test_cli_accepts_legacy_doc_input_when_converter_available(self) -> None:
        if not FORMATTER.shutil.which("textutil") and not (
            FORMATTER.shutil.which("soffice") or FORMATTER.shutil.which("libreoffice")
        ):
            self.skipTest("No .doc conversion tool available.")

        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            docx_path = tmpdir_path / "input.docx"
            doc_path = tmpdir_path / "input.doc"
            output_path = tmpdir_path / "output.docx"

            document = Document()
            document.add_paragraph("关于开展AI办公培训的通知")
            document.add_paragraph("各位老师：")
            document.add_paragraph("请大家按要求完成报名。")
            document.save(docx_path)

            subprocess.run(
                ["textutil", "-convert", "doc", "-output", str(doc_path), str(docx_path)],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
            )
            subprocess.run(
                [sys.executable, str(SCRIPT_PATH), str(doc_path), str(output_path)],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
            )

            formatted = Document(output_path)
            self.assertEqual(cjk_font(formatted.paragraphs[0].runs[0]), "方正小标宋_GBK")


if __name__ == "__main__":
    unittest.main()
